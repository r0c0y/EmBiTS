import os, uuid
from datetime import datetime
from fastapi import FastAPI, File, UploadFile, Form, HTTPException, Request, Depends, Security, BackgroundTasks
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import HTMLResponse, PlainTextResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from database import get_db_connection, init_db
from search import execute_trace
from audit_ledger import log_audit, verify_audit_ledger
from lineage import create_edge, delete_edge, auto_edges
from ingestion import ingest
from registry import router as registry_router
from admin import router as admin_router
from pydantic import BaseModel

API_KEY = os.environ.get("SCL_API_KEY", "")
security = HTTPBearer(auto_error=False)

async def require_auth(creds: HTTPAuthorizationCredentials = Security(security)):
    if not API_KEY:
        return
    if not creds or creds.credentials != API_KEY:
        raise HTTPException(401, "Invalid or missing API key")

app = FastAPI(title="ScribeLink")
BASE_DIR = os.path.dirname(__file__)
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")
app.include_router(registry_router)
app.include_router(admin_router)
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))

@app.on_event("startup")
def startup():
    init_db()
    auto_edges()
    import threading
    from search import build_existing_hierarchical_summaries
    threading.Thread(target=build_existing_hierarchical_summaries, daemon=True).start()

@app.get("/", response_class=HTMLResponse)
async def dashboard(request: Request):
    resp = templates.TemplateResponse(request, "index.html")
    resp.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    resp.headers["Pragma"] = "no-cache"
    return resp

@app.get("/api/documents")
async def get_documents(project: str = None):
    conn = get_db_connection(); c = conn.cursor()
    if project:
        projs = [p.strip() for p in project.split(",") if p.strip()]
        if len(projs) == 1:
            c.execute("SELECT id, title FROM meetings WHERE project_id = ? ORDER BY title ASC", (projs[0],))
        elif len(projs) > 1:
            placeholders = ",".join("?" for _ in projs)
            c.execute(f"SELECT id, title FROM meetings WHERE project_id IN ({placeholders}) ORDER BY title ASC", projs)
        else:
            c.execute("SELECT id, title FROM meetings ORDER BY title ASC")
    else:
        c.execute("SELECT id, title FROM meetings ORDER BY title ASC")
    rows = c.fetchall(); conn.close()
    return {"documents": [dict(r) for r in rows]}

@app.get("/health")
async def health():
    conn = get_db_connection(); c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM meetings")
    doc_count = c.fetchone()[0]
    c.execute("SELECT COUNT(*) FROM chunks")
    chunk_count = c.fetchone()[0]
    conn.close()
    return {"status": "ok", "documents": doc_count, "chunks": chunk_count}

@app.get("/api/projects")
async def get_projects():
    conn = get_db_connection(); c = conn.cursor()
    c.execute("SELECT DISTINCT project_id FROM meetings WHERE project_id IS NOT NULL ORDER BY project_id;")
    projects = [{"id": r["project_id"], "name": r["project_id"].replace("_", " "), "description": r["project_id"]} for r in c.fetchall()]
    c.execute("SELECT id, title, date, lot_id, project_id FROM meetings;")
    docs = [dict(r) for r in c.fetchall()]
    lots = list({d["lot_id"]: {"id": d["lot_id"], "name": d["lot_id"].replace("-"," "), "date": d["date"]} for d in docs}.values())
    conn.close()
    return {"projects": projects, "lots": lots}

@app.get("/api/audit_logs")
async def get_audits():
    conn = get_db_connection(); c = conn.cursor()
    c.execute("SELECT timestamp, username, department, action_type, details, current_hash FROM audit_logs ORDER BY id DESC LIMIT 50;")
    logs = [dict(r) for r in c.fetchall()]; conn.close()
    return {"logs": logs}

@app.post("/api/auth/signup")
async def signup(email: str = Form(...), name: str = Form(...), password: str = Form(...)):
    import hashlib, secrets
    email = email.strip().lower()
    name = name.strip()
    if not email or not name or not password:
        raise HTTPException(400, "All fields are required")
    
    conn = get_db_connection(); c = conn.cursor()
    c.execute("SELECT email FROM users WHERE email = ?", (email,))
    if c.fetchone():
        conn.close()
        raise HTTPException(400, "User already exists with this email")
        
    salt = secrets.token_hex(16)
    password_hash = hashlib.sha256((password + salt).encode("utf-8")).hexdigest()
    
    # First user is admin, otherwise standard user
    c.execute("SELECT COUNT(*) FROM users")
    count = c.fetchone()[0]
    role = "admin" if count == 0 else "user"
    
    c.execute("INSERT INTO users VALUES (?, ?, ?, ?, ?)", (email, name, password_hash, salt, role))
    conn.commit()
    conn.close()
    
    log_audit(name, role, "SIGNUP", f"User registered: {email}")
    return {"status": "success", "user": {"email": email, "name": name, "role": role}}

@app.post("/api/auth/signin")
async def signin(email: str = Form(...), password: str = Form(...)):
    import hashlib
    email = email.strip().lower()
    if not email or not password:
        raise HTTPException(400, "Email and password are required")
        
    conn = get_db_connection(); c = conn.cursor()
    c.execute("SELECT email, name, password_hash, salt, role FROM users WHERE email = ?", (email,))
    row = c.fetchone()
    conn.close()
    
    if not row:
        raise HTTPException(401, "Invalid email or password")
        
    h = hashlib.sha256((password + row["salt"]).encode("utf-8")).hexdigest()
    if h != row["password_hash"]:
        raise HTTPException(401, "Invalid email or password")
        
    log_audit(row["name"], row["role"], "SIGNIN", f"User logged in: {email}")
    return {"status": "success", "user": {"email": row["email"], "name": row["name"], "role": row["role"]}}

@app.post("/api/search")
async def search(query: str = Form(...), project: str = Form(None), document: str = Form(None), date_from: str = Form(None), date_to: str = Form(None), user: str = Form("Unknown"), user_dept: str = Form("General"), auth=Depends(require_auth)):
    if not query.strip(): raise HTTPException(400, "Query cannot be empty.")
    log_audit(user, user_dept, "QUERY", f"Searched: '{query}' (Project: {project}, Doc: {document})")
    try: return execute_trace(query, project=project, doc_id=document, date_from=date_from, date_to=date_to)
    except Exception as e: raise HTTPException(500, str(e))

class TranslationRequest(BaseModel):
    text: str
    target_lang: str

@app.post("/api/translate")
async def translate(req: TranslationRequest, auth=Depends(require_auth)):
    if not req.text.strip():
        return {"status": "success", "translated_text": ""}
    if req.target_lang not in ["Hindi", "Punjabi"]:
        raise HTTPException(400, "Unsupported target language. Supported: Hindi, Punjabi")
    
    from search import translate_text
    translated = translate_text(req.text, req.target_lang)
    if not translated:
         raise HTTPException(500, "Translation failed.")
    return {"status": "success", "translated_text": translated}

@app.post("/api/upload")
async def upload_document(file: UploadFile = File(...), project_id: str = Form("Unknown"), user: str = Form("Unknown"), user_dept: str = Form("General"), background_tasks: BackgroundTasks = None, auth=Depends(require_auth)):
    allowed = {".txt",".pdf",".docx",".xlsx",".xls",".pptx",".png",".jpg",".jpeg",".tiff",".bmp",".csv",".md",".json"}
    ext = "." + file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in allowed:
        raise HTTPException(400, f"Unsupported format. Allowed: {', '.join(sorted(allowed))}")
    raw = await file.read()
    size_kb = len(raw) / 1024
    size_label = f"{size_kb:.1f} KB" if size_kb < 1024 else f"{size_kb/1024:.2f} MB"
    import hashlib, io
    chash = hashlib.sha256(raw).hexdigest()
    conn = get_db_connection(); c = conn.cursor()
    c.execute("SELECT id FROM meetings WHERE project_id = ? AND (file_path = ? OR content_hash = ?)", (project_id, file.filename, chash))
    row = c.fetchone()
    if row:
        conn.close()
        return {"status": "duplicate", "message": "Duplicate document in this project.", "document_id": row[0]}
    try:
        result = ingest(io.BytesIO(raw), file.filename, size=len(raw), base_dir=BASE_DIR)
    except ValueError as e:
        conn.close()
        raise HTTPException(400, str(e))
    doc_id = f"DOC-{uuid.uuid4().hex[:8].upper()}"
    doc_date = result["detected_date"] or datetime.now().strftime("%Y-%m-%d")
    lot_id = f"LOT-{doc_date[:7]}-{doc_id[-4:]}"
    safe_title = "".join(c for c in file.filename.rsplit(".", 1)[0] if c.isalnum() or c in " _-").strip() or "Untitled"

    # Insert row first so we have doc_id for structured OCR storage
    c.execute("INSERT INTO meetings VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (doc_id, safe_title, doc_date, lot_id, project_id, file.filename, len(raw),
         result["text"], result["source_type"],
         chash, datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
         result.get("ocr_quality", "auto"), 0,
         None, None, result.get("ocr_engine"), result.get("page_count", 1), user))

    # Store structured OCR sidecar files with real doc_id
    ocr_result_obj = result.get("ocr_result")
    if ocr_result_obj is not None:
        try:
            from ocr_storage import store_ocr_result
            stored = store_ocr_result(
                base_dir=BASE_DIR,
                doc_id=doc_id,
                filename=file.filename,
                ocr_result=ocr_result_obj,
                engine_name=result.get("ocr_engine", "unknown"),
                page_count=result.get("page_count", 1),
                detected_date=doc_date,
            )
            c.execute("UPDATE meetings SET ocr_json=?, ocr_markdown=?, ocr_engine=?, page_count=? WHERE id=?",
                (stored.get("ocr_json"), stored.get("ocr_markdown"), stored.get("ocr_engine"), stored.get("ocr_page_count"), doc_id))
        except Exception:
            pass

    chunks = __import__("ingestion").chunk_text(result["text"])
    for i, ch in enumerate(chunks):
        c.execute("INSERT INTO chunks VALUES (?,?,?,?)", (f"{doc_id}-{i}", doc_id, i, ch))
    conn.commit(); conn.close()

    # Persist original for preview
    try:
        dest_dir = os.path.join(BASE_DIR, "static", "originals")
        os.makedirs(dest_dir, exist_ok=True)
        dest = os.path.join(dest_dir, f"{doc_id}_{file.filename}")
        with open(dest, "wb") as f:
            f.write(raw)
    except Exception:
        pass

    log_audit(user, user_dept, "UPLOAD", f"Uploaded: {file.filename} ({size_label}, {len(chunks)} chunks, engine={result.get('ocr_engine','native')})")
    auto_edges()
    if background_tasks:
        from search import generate_hierarchical_tree_task
        background_tasks.add_task(generate_hierarchical_tree_task, doc_id, project_id, result["text"], safe_title)
    return {"status": "success", "filename": file.filename, "document_id": doc_id, "chunks": len(chunks)}

@app.get("/api/document/{doc_id}")
async def get_document(doc_id: str):
    conn = get_db_connection(); c = conn.cursor()
    c.execute("SELECT id, title, date, lot_id, project_id, file_path, file_size_bytes, transcript_text, source_type, created_at, ocr_quality, corrections_count, ocr_json, ocr_markdown, ocr_engine, page_count FROM meetings WHERE id=?;", (doc_id,))
    row = c.fetchone(); conn.close()
    if not row: raise HTTPException(404, "Document not found")
    return dict(row)

def docx_to_html(docx_path):
    from docx import Document
    doc = Document(docx_path)
    html_parts = []
    
    html_parts.append("""
    <html>
    <head>
    <meta charset="utf-8">
    <style>
        body {
            font-family: 'Outfit', -apple-system, sans-serif;
            background-color: #f1f5f9;
            margin: 0;
            padding: 20px;
        }
        .page {
            background: #ffffff;
            max-width: 800px;
            margin: 30px auto;
            padding: 60px 80px;
            box-shadow: 0 4px 12px rgba(0, 0, 0, 0.08);
            border-radius: 4px;
            box-sizing: border-box;
            min-height: 29.7cm;
            color: #1e293b;
            line-height: 1.6;
        }
        h1, h2, h3, h4 {
            color: #0f172a;
            margin-top: 1.5em;
            margin-bottom: 0.5em;
        }
        p {
            margin-bottom: 1em;
            text-align: justify;
        }
        table {
            width: 100%;
            border-collapse: collapse;
            margin: 1.5em 0;
        }
        th, td {
            border: 1px solid #cbd5e1;
            padding: 10px;
            text-align: left;
        }
        th {
            background-color: #f1f5f9;
            font-weight: 600;
        }
    </style>
    </head>
    <body>
    <div class="page">
    """)
    
    for element in doc.element.body:
        if element.tag.endswith('p'):
            from docx.text.paragraph import Paragraph
            p = Paragraph(element, doc)
            text = p.text.strip()
            if not text:
                continue
            
            if p.style.name.startswith('Heading 1'):
                html_parts.append(f"<h1>{text}</h1>")
            elif p.style.name.startswith('Heading 2'):
                html_parts.append(f"<h2>{text}</h2>")
            elif p.style.name.startswith('Heading 3'):
                html_parts.append(f"<h3>{text}</h3>")
            else:
                html_parts.append(f"<p>{text}</p>")
                
        elif element.tag.endswith('tbl'):
            from docx.table import Table
            t = Table(element, doc)
            html_parts.append("<table>")
            for r_idx, row in enumerate(t.rows):
                html_parts.append("<tr>")
                for cell in row.cells:
                    tag = "th" if r_idx == 0 else "td"
                    html_parts.append(f"<{tag}>{cell.text.strip()}</{tag}>")
                html_parts.append("</tr>")
            html_parts.append("</table>")
            
    html_parts.append("</div></body></html>")
    return "\n".join(html_parts)

def txt_to_html(txt_path):
    with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
    import html
    escaped_content = html.escape(content)
    return f"""
    <html>
    <head>
    <meta charset="utf-8">
    <style>
        body {{
            font-family: 'Outfit', -apple-system, sans-serif;
            background-color: #f1f5f9;
            margin: 0;
            padding: 20px;
        }}
        .page {{
            background: #ffffff;
            max-width: 800px;
            margin: 30px auto;
            padding: 60px 80px;
            box-shadow: 0 4px 12px rgba(0, 0, 0, 0.08);
            border-radius: 4px;
            box-sizing: border-box;
            min-height: 29.7cm;
            white-space: pre-wrap;
            font-family: monospace;
            font-size: 13.5px;
            color: #1e293b;
            line-height: 1.5;
        }}
    </style>
    </head>
    <body>
    <div class="page">{escaped_content}</div>
    </body>
    </html>
    """

@app.get("/api/preview/{doc_id}")
async def preview_document(doc_id: str):
    import mimetypes
    conn = get_db_connection(); c = conn.cursor()
    c.execute("SELECT file_path, title FROM meetings WHERE id=?;", (doc_id,))
    row = c.fetchone(); conn.close()
    if not row: raise HTTPException(404, "Document not found")
    original_name = row["file_path"] if row["file_path"] else row["title"]
    candidates = [original_name] + [f"{original_name}.{ext}" for ext in ["pdf", "png", "jpg", "jpeg", "docx", "xlsx", "pptx", "txt"]]
    for name in candidates:
        path = os.path.join(BASE_DIR, "static", "originals", f"{doc_id}_{name}")
        if os.path.exists(path):
            if name.lower().endswith(".docx"):
                try:
                    html_content = docx_to_html(path)
                    from fastapi.responses import HTMLResponse
                    return HTMLResponse(content=html_content)
                except Exception as e:
                    pass
            elif name.lower().endswith(".txt"):
                try:
                    html_content = txt_to_html(path)
                    from fastapi.responses import HTMLResponse
                    return HTMLResponse(content=html_content)
                except Exception as e:
                    pass
            mime = mimetypes.guess_type(path)[0] or "application/octet-stream"
            return FileResponse(path, media_type=mime)
    raise HTTPException(404, "Original file not found")

@app.get("/api/download/{doc_id}")
async def download_document(doc_id: str):
    import mimetypes
    conn = get_db_connection(); c = conn.cursor()
    c.execute("SELECT file_path, title FROM meetings WHERE id=?;", (doc_id,))
    row = c.fetchone(); conn.close()
    if not row: raise HTTPException(404, "Document not found")
    original_name = row["file_path"] if row["file_path"] else row["title"]
    candidates = [original_name] + [f"{original_name}.{ext}" for ext in ["pdf", "png", "jpg", "jpeg", "docx", "xlsx", "pptx", "txt"]]
    for name in candidates:
        path = os.path.join(BASE_DIR, "static", "originals", f"{doc_id}_{name}")
        if os.path.exists(path):
            mime = mimetypes.guess_type(path)[0] or "application/octet-stream"
            return FileResponse(path, media_type=mime, filename=name, headers={"Content-Disposition": f'attachment; filename="{name}"'})
    raise HTTPException(404, "Original file not found")

@app.post("/api/lineage/create")
async def api_create_edge(from_node_id: str = Form(...), to_node_id: str = Form(...), relation_type: str = Form("followed_by"), rationale: str = Form(""), user: str = Form("Unknown"), user_dept: str = Form("General"), auth=Depends(require_auth)):
    try: return create_edge(from_node_id, to_node_id, relation_type, rationale, user, user_dept)
    except Exception as e: raise HTTPException(400, str(e))

@app.post("/api/lineage/delete")
async def api_delete_edge(from_node_id: str = Form(...), to_node_id: str = Form(...), user: str = Form("Unknown"), user_dept: str = Form("General"), auth=Depends(require_auth)):
    try: return delete_edge(from_node_id, to_node_id, user, user_dept)
    except Exception as e: raise HTTPException(400, str(e))

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
